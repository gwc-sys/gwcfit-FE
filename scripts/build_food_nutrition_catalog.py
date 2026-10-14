"""Build the app food catalog from All_Foods_Nutrition_Per_100g.xlsx."""
from __future__ import annotations

import json
import sys
from pathlib import Path
from openpyxl import load_workbook

SOURCE_FILE = Path.home() / "Downloads" / "All_Foods_Nutrition_Per_100g.xlsx"
OUTPUT_FILE = Path(__file__).resolve().parents[1] / "src" / "data" / "foodNutritionCatalog.ts"
EXPECTED_HEADERS = ["Food ID", "Food Name", "Main Category", "Subcategory", "Primary Benefit", "Calories (kcal/100 g)", "Protein (g/100 g)", "Carbohydrates (g/100 g)", "Fat (g/100 g)", "Nutrition Basis", "Data Note"]

def number(value: object, row_number: int, label: str) -> float:
    try:
        result = round(float(value), 2)
    except (TypeError, ValueError) as error:
        raise ValueError(f"All Foods row {row_number}: invalid {label}: {value!r}") from error
    if result < 0:
        raise ValueError(f"All Foods row {row_number}: negative {label}: {result}")
    return result

def main() -> None:
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(f"Missing workbook: {SOURCE_FILE}")
    sheet = load_workbook(SOURCE_FILE, read_only=True, data_only=True)["All Foods"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"Unexpected headers: {headers!r}")
    catalog: list[dict[str, object]] = []
    seen_ids: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        food_id, name = str(row[0] or "").strip(), str(row[1] or "").strip()
        if not food_id and not name:
            continue
        if not food_id or not name or food_id in seen_ids:
            raise ValueError(f"All Foods row {row_number}: invalid or duplicate Food ID / Food Name")
        seen_ids.add(food_id)
        main_category = str(row[2] or "Uncategorized").strip()
        subcategory = str(row[3] or "").strip()
        catalog.append({
            "id": food_id, "name": name, "mainCategory": main_category, "subcategory": subcategory,
            "primaryBenefit": str(row[4] or "").strip(), "calories": number(row[5], row_number, "calories"),
            "proteinG": number(row[6], row_number, "protein"), "carbsG": number(row[7], row_number, "carbohydrates"),
            "fatG": number(row[8], row_number, "fat"), "fiberG": 0, "sugarG": 0,
            "nutritionBasis": str(row[9] or "Per 100 g").strip(), "dataNote": str(row[10] or "").strip(),
            "categories": [value for value in (main_category, subcategory) if value], "sources": [SOURCE_FILE.name],
        })
    catalog.sort(key=lambda item: str(item["name"]).casefold())
    payload = json.dumps(catalog, indent=2, ensure_ascii=False)
    output = f'''// Generated from All_Foods_Nutrition_Per_100g.xlsx. Nutrition values are per 100 grams.
export type FoodNutritionItem = {{
  id: string; name: string; mainCategory: string; subcategory: string; primaryBenefit: string;
  calories: number; proteinG: number; carbsG: number; fatG: number; fiberG: number; sugarG: number;
  nutritionBasis: string; dataNote: string; categories: string[]; sources: string[];
}};
export const foodNutritionCatalog: FoodNutritionItem[] = {payload};
export const foodNutritionCategories = Array.from(new Set(foodNutritionCatalog.map((food) => food.mainCategory))).sort((a, b) => a.localeCompare(b));
'''
    if "--check" in sys.argv:
        print(f"Validated {len(catalog)} foods from {SOURCE_FILE.name}")
        return
    OUTPUT_FILE.write_text(output, encoding="utf-8")
    print(f"Wrote {len(catalog)} foods to {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
